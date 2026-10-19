"""Cálculos del tablero de gerencia (mano de obra AIRPLAN T1 JMC).

Este archivo no usa Streamlit: solo lee archivos y calcula. La página
pages/1_Gerencia.py lo llama y dibuja los resultados.

Fuentes (todas en Google Drive, compartidas como "Cualquier persona con el enlace"):
  - BD_Nomina_MO ............ nómina liquidada SINCO, mes a mes (pestañas Nomina_SINCO y Ajustes_Cargo)
  - Parametros_Alertas ...... Proyeccion_A (Proyección de Gastos Generales) y MO_proyectada (SINCO)
  - Curva S ................. AEROPUERTO_JMC_OT1 - datos de mediciones y curva S (hoja AVANCE SEMANA A SEMANA)
  - Cuadro de hitos ......... cuadro.resumen.pagos.hitos (hojas Int..., Item... y PPTO)
  - Base_Tablero_Gerencia ... APU de MO, contratos de MO, presupuesto de la A y umbrales
  - Maestro de Personal ..... salarios (para el % de horas extra)
  - Registros de la app ..... horas diarias (para estimar los meses que aún no tienen nómina)
"""
import io
import os
import re
import unicodedata
import urllib.request

import numpy as np
import pandas as pd

FRENTES_GUIAS = {"15", "16", "17"}
OBRA = [f"{i:02d}" for i in range(1, 15)]
FR_TODOS = [f"{i:02d}" for i in range(1, 18)]
INTERV_A_FRENTE = {"3": "01", "3A": "02", "3B": "03", "4": "04", "5": "05", "5A": "06", "5B": "07", "6": "08", "7": "09",
                   "8": "10", "9": "11", "10": "12", "11": "13", "13": "14", "14": "15", "15": "16", "16": "17"}
MES_CORTO = {1: "Ene", 2: "Feb", 3: "Mar", 4: "Abr", 5: "May", 6: "Jun", 7: "Jul", 8: "Ago", 9: "Sep", 10: "Oct", 11: "Nov", 12: "Dic"}


# ------------------------------------------------------------------ utilidades
def normalizar_nombre(nombre):
    if not isinstance(nombre, str):
        return ""
    n = unicodedata.normalize("NFKD", nombre.strip().upper()).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", n)


def clave_persona(n):
    """Mismo nombre aunque venga 'APELLIDOS NOMBRES' (SINCO) o 'NOMBRES APELLIDOS' (app)."""
    return " ".join(sorted(re.sub(r"[^A-Z ]", "", normalizar_nombre(n)).split()))


def grupo_cargo(c):
    c = str(c).lower()
    if "maestro" in c:
        return "Maestros"
    if "oficial" in c:
        return "Oficiales"
    if "ayudante" in c and "sst" not in c:
        return "Ayudantes"
    if "sin confirmar" in c:
        return "Ayudantes"
    return None


def id_drive(url):
    m = re.search(r"(?:/d/|id=)([\w-]{20,})", str(url))
    return m.group(1) if m else None


def descargar(url, nombre_local=None, html_ok=False):
    """Descarga un archivo de Drive (xlsx subido o Google Sheet) y devuelve los bytes.
    Para pruebas sin internet: variable de entorno GERENCIA_LOCAL_<NOMBRE> con la ruta del archivo."""
    if nombre_local:
        ruta = os.environ.get("GERENCIA_LOCAL_" + nombre_local.upper())
        if ruta:
            with open(ruta, "rb") as f:
                return f.read()
    fid = id_drive(url)
    candidatos = [url] if not fid else [f"https://docs.google.com/spreadsheets/d/{fid}/export?format=xlsx",
                                         f"https://drive.google.com/uc?export=download&id={fid}",
                                         f"https://drive.usercontent.google.com/download?id={fid}&export=download&confirm=t"]
    for u in candidatos:
        try:
            with urllib.request.urlopen(u, timeout=180) as r:
                b = r.read()
        except Exception:
            continue
        if b[:2] == b"PK" or (html_ok and b.lstrip()[:1] == b"<" and b"<table" in b[:200000].lower()):
            return b
    raise ValueError(f"No se pudo descargar {nombre_local or url}. Revise que esté compartido como 'Cualquier persona con el enlace'.")


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


# ------------------------------------------------------------------ lectores
def leer_nomina(b):
    x = pd.read_excel(io.BytesIO(b), sheet_name=None, dtype=str)
    n = x["Nomina_SINCO"]
    aj = x.get("Ajustes_Cargo", pd.DataFrame(columns=["Nombre", "Grupo"]))
    return n, aj


def leer_parametros(b):
    x = pd.read_excel(io.BytesIO(b), sheet_name=None, dtype={"Código": str, "Cód. frente": str})
    pa = x["Proyeccion_A"].copy()
    for c in ["Inicio", "Fin"]:
        pa[c] = pd.to_datetime(pa[c], errors="coerce")
    mop = x["MO_proyectada"].copy()
    mop["Código"] = mop["Código"].astype(str).str.strip()
    return pa, mop


def leer_base(b):
    x = pd.read_excel(io.BytesIO(b), sheet_name=None, dtype={"Código": str, "Frente": str, "Cód. frente": str})
    par = dict(zip(x["Parametros"]["Parámetro"].astype(str), x["Parametros"]["Valor"]))
    for k, v in list(par.items()):
        try:
            par[k] = float(v)
        except (TypeError, ValueError):
            par[k] = str(v)
    fr = x["Frentes"].copy()
    fr["Cód. frente"] = fr["Cód. frente"].astype(str).str.zfill(2)
    return dict(apu=x["APU_MO"], contratos=x["Contratos_MO"], ppto_a=x["Presupuesto_A"], frentes=fr, par=par)


def leer_curva(b):
    import openpyxl
    rows = list(openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)["AVANCE SEMANA A SEMANA"].iter_rows(values_only=True))
    h2 = rows[1]
    semanas, c = [], 10
    while c < len(h2) and h2[c] is not None:
        semanas.append((pd.Timestamp(h2[c]).normalize(), c))
        c += 8
    regs = []
    for r in rows[5:]:
        if r[1] is None:
            continue
        wbs = str(r[1]).strip()
        if wbs != "0" and wbs.count(".") != 0:
            continue
        nombre = str(r[2] if r[2] is not None else (r[3] or "")).strip()
        for f, c in semanas:
            regs.append((wbs, nombre, num(r[5]), f, num(r[c + 3]), num(r[c + 7])))
    P = pd.DataFrame(regs, columns=["wbs", "nombre", "costo", "semana", "PV", "EV"])
    P["fr"] = np.where(P.wbs == "0", "TOTAL", P.nombre.str.extract(r"^\((\w+)\)")[0].map(INTERV_A_FRENTE))
    return P


def leer_hitos(b):
    """Hitos con lo programado y cobrado por acta, sus ítems y el valor de cada ítem del presupuesto."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    H, I = [], []
    for ws in wb.worksheets:
        if not ws.title.startswith(("Int", "Item")):
            continue
        cur, frh = None, None
        for i, r in enumerate(ws.iter_rows(min_row=1, max_col=78, values_only=True)):
            if i < 4:
                continue
            r = list(r) + [None] * (78 - len(r))
            if r[0] is None and r[1] is not None and r[2] is not None and r[3] and str(r[3]).startswith("("):
                frh = str(r[2]).strip()
                continue
            if r[0] not in (None, 0, "0") and r[3] is not None:
                cur = dict(hid=len(H), fr=frh, hito=str(r[0]), nombre=r[3], valor=num(r[10]))
                for k in range(1, 27):
                    cur[f"p{k}"] = num(r[20 + 2 * (k - 1)])
                    cur[f"c{k}"] = num(r[21 + 2 * (k - 1)])
                H.append(cur)
                continue
            if cur is not None and r[2] is not None and r[0] is None and r[10] is not None:
                I.append(dict(hid=cur["hid"], cod=str(r[2]).strip(), valor_item=num(r[10])))
    H, I = pd.DataFrame(H), pd.DataFrame(I)
    rows = list(wb["PPTO"].iter_rows(min_row=4, values_only=True))
    pp = pd.DataFrame([(str(r[0]).strip(), r[2], num(r[6])) for r in rows if r[0] is not None], columns=["cod", "desc", "valor"])
    pp = pp[pp.cod.str.match(r"^\d\d")]
    cods = set(pp.cod)
    pp["hoja"] = [not any(o.startswith(c + ".") for o in cods) for c in pp.cod]
    venta = pp[pp.hoja].drop_duplicates("cod").set_index("cod")
    return H, I, venta


def leer_maestro(b):
    t = pd.read_excel(io.BytesIO(b), sheet_name="Maestro Personal", header=None)
    h = t[t[0] == "Documento"].index[0]
    t.columns = t.iloc[h]
    t = t.iloc[h + 1:].dropna(subset=["Documento"]).copy()
    t["Documento"] = t["Documento"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    t["SalarioMensual"] = pd.to_numeric(t["SalarioMensual"], errors="coerce")
    t["clave"] = t["NombreCompleto"].map(clave_persona)
    return t


def leer_excluidos(b):
    """Pestaña Excluidos de Parametros_Alertas: [(clave, desde)]. Si no existe, lista vacía."""
    try:
        ex = pd.read_excel(io.BytesIO(b), sheet_name="Excluidos", dtype=str)
    except Exception:
        return []
    ex = ex.dropna(subset=[ex.columns[0]])
    desde = ex["Desde"] if "Desde" in ex.columns else [None] * len(ex)
    return [(clave_persona(n), pd.to_datetime(d, errors="coerce")) for n, d in zip(ex.iloc[:, 0], desde)]


def quitar_excluidos(df, lista, col_nombre="nombre", col_fecha="fecha"):
    if not lista or df is None or not len(df):
        return df
    clave = df[col_nombre].map(clave_persona)
    quitar = pd.Series(False, index=df.index)
    for k, d in lista:
        m_ = clave == k
        if pd.notna(d):
            m_ &= df[col_fecha] >= d
        quitar |= m_
    return df[~quitar].copy()


# ------------------------------------------------------------------ nómina
def base_nomina(n, aj, estimado_app=None):
    """Una fila por persona × mes × código con costo, grupo e imputación (maestros y GG a la A,
    oficiales y ayudantes al APU de la actividad, guías a su ítem)."""
    n = n.dropna(subset=["Cedula"]).copy()
    for c in ["Valor", "Porcentaje", "Anio", "Mes"]:
        n[c] = pd.to_numeric(n[c], errors="coerce")
    n = n.dropna(subset=["Valor", "Anio", "Mes"])
    n["Porcentaje"] = n["Porcentaje"].fillna(1)
    R = pd.DataFrame({"mes": [pd.Period(year=int(a), month=int(m), freq="M") for a, m in zip(n.Anio, n.Mes)],
                      "fuente": "Nómina SINCO", "cedula": n.Cedula.astype(str).str.replace(r"\.0$", "", regex=True).str.strip(),
                      "nombre": n["Nombre Empleado"].astype(str).str.strip(), "cargo": n.Cargo,
                      "cod": n["No Item Ppto"].fillna("").astype(str).str.replace(r"\s+", "", regex=True),
                      "costo": n.Valor * n.Porcentaje})
    meses_nom = sorted(R.mes.unique())
    if estimado_app is not None and len(estimado_app):
        R = pd.concat([R, estimado_app[estimado_app.mes > max(meses_nom)]], ignore_index=True)
    R["clave"] = R.nombre.map(clave_persona)
    R["grupo"] = R.cargo.map(grupo_cargo)
    if len(aj) and {"Nombre", "Grupo"}.issubset(aj.columns):
        # columna opcional "Desde" (año-mes, ej. 2026-09): el ajuste aplica desde ese mes
        desde = aj["Desde"] if "Desde" in aj.columns else pd.Series([None] * len(aj), index=aj.index)
        for a, g_, d_ in zip(aj.Nombre, aj.Grupo, desde):
            if not (isinstance(a, str) and isinstance(g_, str)):
                continue
            sel = R.clave == clave_persona(a)
            if isinstance(d_, str) and d_.strip() and d_.strip().lower() != "nan":
                try:
                    sel &= R.mes >= pd.Period(d_.strip()[:7], "M")
                except Exception:
                    pass
            R.loc[sel, "grupo"] = g_.strip()
    R = R[R.grupo.isin(["Maestros", "Oficiales", "Ayudantes"])].copy()
    cu = R.cod.str.upper()
    R["fr"] = R.cod.str.extract(r"^(\d\d)\.")[0].fillna(R.cod.str.extract(r"^(?:NP|SUP)-(\d\d)", flags=re.I)[0])
    destino = np.select([R.cod == "", cu.str.endswith("GG"), R.fr.isin(FRENTES_GUIAS), cu.str.startswith("SUP")],
                        ["Sin código", "GG", "Guías", "Supervisión"], "Ítem")
    R["imputa"] = np.where(R.grupo == "Maestros", "A", np.where(np.isin(destino, ["GG", "Supervisión", "Sin código"]), "A",
                           np.where(destino == "Guías", "Guías", "APU")))
    return R, meses_nom


def estimar_con_app(mo, meses_nom):
    """Meses sin nómina: horas de la app × tarifa × factor de recargo (columna costo_calculado de la app)."""
    if mo is None or not len(mo):
        return None
    ap = mo.dropna(subset=["fecha"]).copy()
    ap["mes"] = ap.fecha.dt.to_period("M")
    ap = ap[ap.mes > max(meses_nom)]
    if not len(ap):
        return None
    return pd.DataFrame({"mes": ap.mes, "fuente": "App (estimado)", "cedula": "", "nombre": ap.nombre.astype(str),
                         "cargo": ap.cargo, "cod": ap.codact.astype(str).str.strip(), "costo": ap.costo_calculado.fillna(0),
                         "fecha_max": ap.fecha.max()})


# ------------------------------------------------------------------ cálculo principal
def calcular(R, meses_nom, pa, mop, P, H, I, venta_items, base, maestro=None, hoy=None):
    hoy = pd.Timestamp(hoy or pd.Timestamp.today()).normalize()
    par = base["par"]
    out = {}
    fr_nom = base["frentes"].set_index("Cód. frente")
    nombre_fr = fr_nom["Frente"].to_dict()
    residente = fr_nom["Residente"].to_dict()
    # Las comparaciones usan solo la nómina SINCO (cifra verificada); el estimado de la app se muestra aparte.
    Rall = R
    R = Rall[Rall.fuente == "Nómina SINCO"]
    est = Rall[Rall.fuente == "App (estimado)"]
    ult_mes = max(meses_nom)
    corte_real = ult_mes.end_time.normalize()
    out["cortes"] = dict(nomina=ult_mes, real_hasta=corte_real, estimado=len(est) > 0,
                         estimado_hasta=(est.fecha_max.max() if "fecha_max" in est and len(est) else None))
    out["estimado_mes"] = est.costo.sum()

    # ---------- cobro (hitos)
    actas_cob = [k for k in range(1, 27) if H[f"c{k}"].sum() > 0]
    ult_acta = max(actas_cob)
    H = H.copy()
    H["cob"] = H[[f"c{k}" for k in range(1, ult_acta + 1)]].sum(axis=1)
    H["prog"] = H[[f"p{k}" for k in range(1, ult_acta + 1)]].sum(axis=1)
    H["prog_sig"] = H[f"p{ult_acta + 1}"] if ult_acta < 26 else 0.0
    out["cortes"]["acta"] = ult_acta
    # cobrado por ítem: valor del ítem en el hito × fracción cobrada del hito; los capítulos se reparten por valor
    hojas = venta_items[venta_items.index != "3"]
    fc = H.set_index("hid").eval("cob / valor").replace([np.inf, -np.inf], 0).fillna(0)
    Ix = I.copy()
    Ix["cob"] = Ix.valor_item * Ix.hid.map(fc).fillna(0)
    filas = []
    for cod, g in Ix.groupby("cod"):
        cob = g.cob.sum()
        if cod in hojas.index:
            filas.append((cod, cob))
        else:
            kids = hojas[hojas.index.str.startswith(cod + ".")]
            tot = kids.valor.sum()
            if tot > 0:
                filas.extend((k, cob * v / tot) for k, v in kids.valor.items())
            else:
                filas.append((cod, cob))
    cob_item = pd.DataFrame(filas, columns=["cod", "cob"]).groupby("cod").cob.sum()
    it = pd.DataFrame({"venta": hojas.valor}).join(cob_item.rename("cob"), how="outer").fillna(0)
    it["pc"] = np.where(it.venta > 0, (it.cob / it.venta).clip(upper=1), 0)
    it["fr"] = it.index.str[:2]

    # ---------- avance (curva S)
    Pf = P[(P.fr != "TOTAL") & P.fr.notna()]
    sem_ok = sorted(s for s in Pf.semana.unique() if s <= hoy + pd.Timedelta(days=6) and Pf[Pf.semana == s].EV.sum() > 0)
    ult_sem = sem_ok[-1]
    ev = Pf.pivot_table(index="semana", columns="fr", values="EV", aggfunc="sum").reindex(columns=FR_TODOS).fillna(0)
    pv = Pf.pivot_table(index="semana", columns="fr", values="PV", aggfunc="sum").reindex(columns=FR_TODOS).fillna(0)
    rezago = int(par.get("dias_rezago_curva", 4))

    def sem_para(fecha):
        c = [s for s in sem_ok if s >= fecha + pd.Timedelta(days=rezago)]
        return c[0] if c else ult_sem

    sem_real = sem_para(corte_real)
    out["cortes"]["curva"] = ult_sem
    out["cortes"]["curva_comparacion"] = sem_real

    # ---------- frentes
    venta_fr = H.groupby("fr").valor.sum().reindex(FR_TODOS).fillna(0)
    mop2 = mop.copy()
    mop2["fr"] = mop2["Código"].str[:2]
    mo_meta = pd.to_numeric(mop2["Proy. mano de obra"], errors="coerce").fillna(0).groupby(mop2.fr).sum().reindex(FR_TODOS).fillna(0)
    ratio = (mo_meta / venta_fr).where(venta_fr > 0, 0)
    real_fr = R[R.imputa != "A"].groupby("fr").costo.sum().reindex(FR_TODOS).fillna(0)
    F = pd.DataFrame(index=FR_TODOS)
    F["frente"] = [nombre_fr.get(k, k) for k in F.index]
    F["residente"] = [residente.get(k, "") for k in F.index]
    F["venta"] = venta_fr
    F["cobrado"] = H.groupby("fr").cob.sum().reindex(FR_TODOS).fillna(0)
    F["programado"] = H.groupby("fr").prog.sum().reindex(FR_TODOS).fillna(0)
    F["prog_siguiente"] = H.groupby("fr").prog_sig.sum().reindex(FR_TODOS).fillna(0)
    F["ejecutado"] = ev.loc[ult_sem]
    F["programado_curva"] = pv.loc[ult_sem]
    F["pct_ejecutado"] = np.where(F.venta > 0, F.ejecutado / F.venta, 0)
    F["pct_cobrado"] = np.where(F.venta > 0, F.cobrado / F.venta, 0)
    F["mo_meta"] = mo_meta
    F["real"] = real_fr
    F["estimado_mes"] = est[est.imputa != "A"].groupby("fr").costo.sum().reindex(FR_TODOS).fillna(0) if len(est) else 0.0
    F["pct_avance_comp"] = np.where(F.venta > 0, ev.loc[sem_real] / F.venta, 0)
    F["meta_ejecutado"] = F.mo_meta * F.pct_avance_comp
    F["pct_gastado"] = np.where(F.meta_ejecutado > 0, F.real / F.meta_ejecutado, np.nan)
    F["cierre"] = np.where((F.pct_avance_comp >= 0.10) & (F.real > 1e6), F.real / F.pct_avance_comp.replace(0, np.nan), np.nan)

    # actividades: nómina contra MO del APU
    apu = base["apu"].set_index("Código")["MO oficiales y ayudantes del APU"]
    nom_it = R[R.imputa == "APU"].groupby("cod").agg(nomina=("costo", "sum"), personas=("clave", "nunique"))
    A = nom_it.join(apu.rename("apu"), how="left").fillna({"apu": 0})
    A["fr"] = A.index.str[:2]
    A["pc"] = it.pc.reindex(A.index).fillna(0)
    A["desc"] = [str(venta_items.desc.get(k, "")) for k in A.index]
    A["tipo"] = np.select([(A.apu <= 0) & (A.nomina > 0), A.nomina > A.apu], ["Sin MO en el APU", "Nómina por encima del APU"], "")
    A["valor"] = np.where(A.tipo == "Nómina por encima del APU", A.nomina - A.apu, np.where(A.tipo == "Sin MO en el APU", A.nomina, 0))
    ap = A[A.tipo == "Nómina por encima del APU"]
    F["valor_sobre_apu"] = ap.groupby("fr").valor.sum().reindex(FR_TODOS).fillna(0)
    F["n_sobre_apu"] = ap.groupby("fr").size().reindex(FR_TODOS).fillna(0)

    # contratos de MO
    C = base["contratos"].copy()
    C["Código"] = C["Código"].astype(str)
    C["fr"] = C.Frente.astype(str).str.zfill(2)
    for c in ["Contratado MO (antes de IVA)", "Otros contratos comparables", "Pagado MO (actas antes de IVA)", "Presupuesto SINCO comparable"]:
        C[c] = pd.to_numeric(C[c], errors="coerce").fillna(0)
    C["nomina"] = C["Código"].map(R[R.imputa != "A"].groupby("cod").costo.sum()).fillna(0)
    C["sobre"] = C["Contratado MO (antes de IVA)"] + C["Otros contratos comparables"] + C.nomina - C["Presupuesto SINCO comparable"]
    C["pc"] = C["Código"].map(it.pc).fillna(0)
    C["sin_cobrar"] = (C["Pagado MO (actas antes de IVA)"] - C["Contratado MO (antes de IVA)"] * C.pc).clip(lower=0)
    F["contratado_sobre"] = C[C.sobre > 0].groupby("fr").sobre.sum().reindex(FR_TODOS).fillna(0)
    F["sub_pagado"] = C.groupby("fr")["Pagado MO (actas antes de IVA)"].sum().reindex(FR_TODOS).fillna(0)
    F["sub_contratado"] = C.groupby("fr")["Contratado MO (antes de IVA)"].sum().reindex(FR_TODOS).fillna(0)
    F["sin_cobrar"] = C.groupby("fr").sin_cobrar.sum().reindex(FR_TODOS).fillna(0)
    out["sin_cobrar_fuera"] = C[~C.fr.isin(FR_TODOS)].sin_cobrar.sum()
    out["sub_pagado_fuera"] = C[~C.fr.isin(FR_TODOS)]["Pagado MO (actas antes de IVA)"].sum()

    # ---------- horas extra
    he = horas_extra(R, pa, maestro, par)
    out.update(he)
    F["he_sobre_40"] = he["he_frente"].reindex(FR_TODOS).fillna(0)

    # ---------- semáforo y qué decidir
    def estado(r):
        if r.real <= 0 and r.contratado_sobre <= 0 and r.sin_cobrar <= 0:
            return ""
        rojo = (r.meta_ejecutado > 0 and r.pct_gastado > par["rojo_pct_presupuesto"]) or r.valor_sobre_apu >= par["rojo_valor"] \
            or r.contratado_sobre >= par["rojo_valor"] or r.sin_cobrar >= par["rojo_pagado_sin_cobrar"]
        if rojo:
            return "Rojo"
        amar = (r.meta_ejecutado > 0 and r.pct_gastado >= par["amarillo_pct_presupuesto"]) or r.valor_sobre_apu >= par["amarillo_valor"] \
            or r.contratado_sobre >= par["amarillo_valor"] or r.sin_cobrar >= par["amarillo_pagado_sin_cobrar"] or r.he_sobre_40 >= par["amarillo_he_sobre_40"]
        return "Amarillo" if amar else "Verde"

    out["frentes"] = F
    out["actividades"] = A[A.tipo != ""].copy()
    out["contratos"] = C
    out["items"] = it

    # ---------- maestros: real contra proyección y presupuesto
    hm = he['he_mes'].tail(3)
    he_real = hm.real.sum() / hm.base.sum() - 1 if len(hm) else 0.30
    out.update(proyeccion_a(R, pa, base, corte_real, he_real))

    # ---------- MO propia mes a mes (nómina contra MO ganada por el avance)
    meses = [m for m in sorted(R.mes.unique()) if m >= pd.Period("2026-04", "M")]
    rm = R[(R.imputa == "APU") & R.fr.isin(OBRA)].pivot_table(index="mes", columns="fr", values="costo", aggfunc="sum").reindex(index=meses, columns=OBRA).fillna(0)
    meta = pd.DataFrame({m: ev.loc[sem_para(min(m.end_time.normalize(), corte_real))].reindex(OBRA).fillna(0) * ratio.reindex(OBRA) for m in meses}).T
    # abril y mayo juntos: la curva S empezó a medir en junio
    if len(meses) > 1 and meses[0].month == 4:
        rm.iloc[1] += rm.iloc[0]
        rm, meta = rm.iloc[1:], meta.iloc[1:]
    dmeta = meta.diff()
    dmeta.iloc[0] = meta.iloc[0]
    T = pd.DataFrame({"nomina": rm.sum(axis=1), "ganada": dmeta.sum(axis=1), "nomina_acum": rm.cumsum().sum(axis=1), "ganada_acum": meta.sum(axis=1)})
    out["tendencia"] = T
    out["tendencia_fr"] = (rm, dmeta)
    F["nomina_ult_mes"] = rm.iloc[-1].reindex(FR_TODOS).fillna(0)
    F["ganada_ult_mes"] = dmeta.iloc[-1].reindex(FR_TODOS).fillna(0)
    F["semaforo"] = F.apply(estado, axis=1)
    F["decidir"] = F.apply(lambda r: decidir(r, par), axis=1)

    # ---------- gasto mensual por grupo
    R2 = Rall.copy()
    R2["serie"] = np.where(R2.imputa == "A", np.where(R2.grupo == "Maestros", "Administración: maestros", "Administración: oficiales y ayudantes"),
                           np.where(R2.imputa == "Guías", "Guías", R2.fr.map(lambda k: nombre_fr.get(k, "Otros"))))
    out["gasto_mes"] = R2.groupby(["mes", "serie", "fuente"]).costo.sum().reset_index()
    out["R"] = Rall
    return out


def decidir(r, par):
    t = []
    if r.meta_ejecutado > 0 and r.pct_gastado > par["rojo_pct_presupuesto"]:
        t.append(f"Nómina propia al {r.pct_gastado:.0%} de su presupuesto por lo ejecutado; al ritmo actual cierra en {m(r.cierre)} contra {m(r.mo_meta)}. Revisar cuadrilla y rendimiento.")
    elif r.meta_ejecutado > 0 and r.pct_gastado >= par["amarillo_pct_presupuesto"]:
        t.append(f"Nómina al límite ({r.pct_gastado:.0%}). Vigilar cada semana.")
    if r.ganada_ult_mes > 1e6 and r.nomina_ult_mes > 1.5 * r.ganada_ult_mes:
        t.append(f"Último mes: {m(r.nomina_ult_mes)} de nómina contra {m(r.ganada_ult_mes)} ganados por el avance.")
    if r.contratado_sobre >= par["amarillo_valor"]:
        t.append(f"{m(r.contratado_sobre)} contratados por encima de SINCO: revisar antes de pagar saldos o tramitar mayor valor.")
    if r.sin_cobrar >= par["amarillo_pagado_sin_cobrar"]:
        t.append(f"{m(r.sin_cobrar)} pagados a subcontratistas por obra sin cobrar: incluir en la próxima acta.")
    if r.ejecutado - r.cobrado > 100e6:
        t.append(f"{m(r.ejecutado - r.cobrado)} ejecutados sin cobrar.")
    if r.valor_sobre_apu >= par["amarillo_valor"]:
        t.append(f"{int(r.n_sobre_apu)} actividades con nómina por encima del APU ({m(r.valor_sobre_apu)}).")
    if r.he_sobre_40 >= par["amarillo_he_sobre_40"]:
        t.append(f"{m(r.he_sobre_40)} pagados sobre el 40 % de horas extra.")
    return " ".join(t)


def m(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "—"
    s = f"{v / 1e6:,.1f}".replace(",", "#").replace(".", ",").replace("#", ".")
    return f"{s} M"


def horas_extra(R, pa, maestro, par):
    """% pagado sobre el básico por persona y mes completo; lo pagado por encima del tope."""
    tope = float(par.get("tope_he", 0.40))
    N = R[R.fuente == "Nómina SINCO"]
    pm = N.groupby(["clave", "mes"]).agg(real=("costo", "sum"), grupo=("grupo", "first"), nombre=("nombre", "first"), cedula=("cedula", "first")).reset_index()
    base = {}
    p = pa.copy()
    p["clave"] = p["Nombre en la app (Maestro de Personal)"].map(lambda x: clave_persona(x) if isinstance(x, str) and x.strip() else None)
    q = p.dropna(subset=["clave"])
    for cl, s_, sub, f_ in zip(q.clave, pd.to_numeric(q["Salario básico"], errors="coerce"), pd.to_numeric(q["Subsidio"], errors="coerce"),
                               pd.to_numeric(q["Factor prestacional"], errors="coerce")):
        if pd.notna(s_) and s_ > 0:
            base[cl] = (s_ + (sub if pd.notna(sub) else 350000)) * (f_ if pd.notna(f_) else 1.4914)
    if maestro is not None:
        ced = dict(zip(maestro.Documento, maestro.SalarioMensual))
        cla = dict(zip(maestro.clave, maestro.SalarioMensual))
        for r in pm.drop_duplicates("clave").itertuples():
            if r.clave not in base:
                s = ced.get(r.cedula, cla.get(r.clave))
                if s and s > 0:
                    base[r.clave] = (s + 350000) * 1.4914
    pm["base"] = pm.clave.map(base)
    meses = sorted(pm.mes.unique())
    idx = {mm: i for i, mm in enumerate(meses)}
    pres = set(zip(pm.clave, pm.mes.map(idx)))
    ult = len(meses) - 1
    pm["completo"] = [((c, idx[mm] - 1) in pres) and (((c, idx[mm] + 1) in pres) or idx[mm] == ult) for c, mm in zip(pm.clave, pm.mes)]
    f = pm[pm.completo & pm.base.notna()].copy()
    f["extra"] = f.real / f.base - 1
    f["pagado_extra"] = f.real - f.base
    f["proyectado_extra"] = f.base * tope
    porm = f.groupby("mes").agg(personas=("real", "size"), real=("real", "sum"), base=("base", "sum"), pagado=("pagado_extra", "sum"), proyectado=("proyectado_extra", "sum"))
    porg = f.groupby(["grupo", "mes"]).apply(lambda d: d.real.sum() / d.base.sum() - 1).unstack()
    red = f[f.extra > tope].copy()
    red["exceso"] = red.real - red.base * (1 + tope)
    rr = R.merge(red[["clave", "mes", "real", "exceso"]], on=["clave", "mes"])
    rr["asig"] = rr.exceso * rr.costo / rr.real
    rr["frx"] = np.where(rr.imputa == "A", "A", rr.fr)
    he_frente = rr.groupby("frx").asig.sum()
    he_act = rr[rr.imputa != "A"].groupby("cod").agg(valor=("asig", "sum"), personas=("clave", "nunique"))
    personas = red.groupby(["nombre", "grupo"]).agg(meses=("mes", lambda s: ", ".join(MES_CORTO[x.month].lower() for x in sorted(s))),
                                                  exceso=("exceso", "sum"), maximo=("extra", "max")).reset_index().sort_values("exceso", ascending=False)
    return dict(he_mes=porm, he_grupo=porg, he_frente=he_frente, he_act=he_act, he_personas=personas, he_tope=tope)


def proyeccion_a(R, pa, base, corte_real, he_real=0.30):
    p = pa[(pa["¿Proyectado en A?"].astype(str).str.strip() == "Sí") & (pa["Sección"].astype(str).str.strip() == "PERSONAL ADMINISTRATIVO")].copy()
    p["g"] = np.where(p["Cargo en la proyección"].astype(str).str.upper().str.contains("MAESTRO"), "Maestros", "Oficiales y ayudantes")
    p["mensual"] = pd.to_numeric(p["Salario total mes proyectado (con HE)"], errors="coerce").fillna(0)
    fin = pd.Timestamp(str(base["par"].get("fin_obra", "2027-04-24")))
    meses = pd.period_range("2026-04", fin.to_period("M"), freq="M")
    filas = []
    for r in p.itertuples():
        if pd.isna(r.Inicio) or pd.isna(r.Fin):
            continue
        for mm in meses:
            a, b = mm.start_time.normalize(), mm.end_time.normalize()
            d = (min(b, r.Fin) - max(a, r.Inicio)).days + 1
            if d > 0:
                filas.append((mm, r.g, r.mensual * d / b.day))
    PM = pd.DataFrame(filas, columns=["mes", "g", "proy"]).pivot_table(index="mes", columns="g", values="proy", aggfunc="sum").reindex(meses).fillna(0)
    RA = R[R.imputa == "A"].copy()
    RA["g"] = np.where(RA.grupo == "Maestros", "Maestros", "Oficiales y ayudantes")
    ra = RA.pivot_table(index="mes", columns="g", values="costo", aggfunc="sum")
    ult = max(R.mes)
    ppto = dict(zip(base["ppto_a"].iloc[:, 0], pd.to_numeric(base["ppto_a"].iloc[:, 1], errors="coerce")))
    real_m = ra["Maestros"].sum() if "Maestros" in ra else 0.0
    proy_hasta = PM.loc[:ult, "Maestros"].sum()
    resto = PM.loc[ult + 1:, "Maestros"].sum()
    corte = ult.end_time.normalize()
    meses_rest = len(pd.period_range(ult + 1, fin.to_period('M') - 1, freq='M')) + fin.day / fin.days_in_month if fin > corte else 0.0
    ritmo = ra["Maestros"].iloc[-1] if "Maestros" in ra and len(ra) else 0.0
    ppto_m = ppto.get("Maestros", np.nan)
    EA = dict(real=real_m, proy_hasta=proy_hasta, proy_resto=resto, proy_total=PM.Maestros.sum(), ppto=ppto_m,
              cierre_plan=real_m + resto, cierre_ritmo=real_m + ritmo * meses_rest, ritmo=ritmo, meses_rest=meses_rest,
              tope_mes=(ppto_m - real_m) / meses_rest if meses_rest > 0 else np.nan,
              cierre_he_real=real_m + resto * (1 + he_real) / (1 + float(base['par'].get('tope_he', 0.40))), he_real=he_real)
    oa_proy = PM.loc[:ult, "Oficiales y ayudantes"].sum() if "Oficiales y ayudantes" in PM else 0.0
    oa_real = ra["Oficiales y ayudantes"].sum() if "Oficiales y ayudantes" in ra else 0.0
    adm = PM.join(ra.add_prefix("real_"))
    return dict(maestros=EA, oa=dict(proy=oa_proy, real_a=oa_real), adm_mes=adm)


# ------------------------------------------------------------------ subcontratistas (exportes SINCO)
_COD = re.compile(r"(\d{2}(?:\.\d{2}){1,5}(?:\.\d{1,2})?(?:GG)?)")


def _tabla_html(b):
    if isinstance(b, (bytes, bytearray)) and b[:2] == b"PK":   # si Drive lo convirtió a Hoja de cálculo
        x = pd.read_excel(io.BytesIO(b), header=None)
        fila = next(i for i in range(min(15, len(x))) if x.iloc[i].astype(str).str.contains("Contrato").any())
        t = x.iloc[fila + 1:].copy()
        t.columns = [str(c) for c in x.iloc[fila]]
        return t.reset_index(drop=True)
    s = b.decode("latin-1") if isinstance(b, (bytes, bytearray)) else b
    t = max(pd.read_html(io.StringIO(s)), key=lambda x: x.size)
    t.columns = [str(c[-1]) if isinstance(c, tuple) else str(c) for c in t.columns]
    return t


def leer_contratos_sinco(b):
    """SINCO > Informe contratos detallado (archivo .xls tal como sale de SINCO)."""
    t = _tabla_html(b).iloc[:, :20]
    t.columns = ["proy", "cod_ctta", "contrato", "grupo", "ctta", "actividad", "insumo", "insumo_desc", "razon", "anio",
                 "f_elab", "f_ini", "v_orig", "v_adic", "v_contrato", "cant", "und", "vu", "vu_iva", "vt_iva"]
    t["contrato"] = pd.to_numeric(t.contrato, errors="coerce")
    t = t[t.contrato.notna()].copy()
    t["contrato"] = t.contrato.astype(int)
    for c in ["grupo", "v_contrato", "cant", "vu", "vt_iva"]:
        t[c] = pd.to_numeric(t[c], errors="coerce")
    t["cod"] = t.actividad.astype(str).str.extract(r"^\s*([\w\.]+)\s*-")[0]
    t["base"] = t.cant.fillna(0) * t.vu.fillna(0)
    t["ctta"] = t.ctta.astype(str).str.strip()
    return t


def leer_actas_sinco(b, C):
    """SINCO > Detallado de actas (archivo .xls tal como sale de SINCO). El ítem sale del grupo del contrato o del concepto."""
    t = _tabla_html(b).iloc[:, :10]
    t.columns = ["concepto", "tipo", "um", "contrato", "acta", "fecha", "cant", "vu", "iva", "vt"]
    t["contrato"] = pd.to_numeric(t.contrato, errors="coerce")
    t = t[t.contrato.notna()].copy()
    t["contrato"] = t.contrato.astype(int)
    for c in ["acta", "cant", "vu", "iva", "vt"]:
        t[c] = pd.to_numeric(t[c], errors="coerce")
    t["fecha"] = pd.to_datetime(t.fecha, dayfirst=True, errors="coerce")
    t["base"] = t.cant.fillna(0) * t.vu.fillna(0)
    t["grupo"] = pd.to_numeric(t.concepto.astype(str).str.extract(r"^\s*(\d+)\s*-")[0], errors="coerce")
    g = C.dropna(subset=["grupo"]).groupby(["contrato", "grupo"]).cod.first()
    t["cod"] = [g.get((c, gr)) if pd.notna(gr) and str(tp).strip() == "Grupo" else None for c, gr, tp in zip(t.contrato, t.grupo, t.tipo)]
    t["cod"] = t.cod.fillna(t.concepto.astype(str).str.extract(_COD)[0])
    return t


def subcontratistas(C, A, mop, items, P, hoy=None):
    """Contratado, pagado, pagado sin cobrar y contratado por encima de la proyección SINCO, por ítem, frente y contratista.
    Ítems directos = los que están en MO_proyectada (proyección SINCO por ítem); el resto se reporta como indirectos."""
    hoy = pd.Timestamp(hoy or pd.Timestamp.today()).normalize()
    m = mop.copy()
    m["Código"] = m["Código"].astype(str).str.strip()
    m = m.drop_duplicates("Código").set_index("Código")
    con = C.groupby("cod").base.sum()
    pag = A.groupby("cod").base.sum()
    it = pd.DataFrame({"contratado": con, "pagado": pag}).fillna(0)
    it = it[(it.contratado > 0) | (it.pagado > 0)]
    it["directo"] = it.index.isin(m.index)
    it["fr"] = it.index.map(m["Frente"].astype(str).str.zfill(2)).fillna("—")
    it["desc"] = it.index.map(m["Descripción"]).fillna("")
    it["proyectado"] = it.index.map(m["Proyectado total"]).fillna(0)
    it["pc"] = it.index.map(items.pc if items is not None and "pc" in items else pd.Series(dtype=float)).fillna(0)
    it["cobrado"] = it.index.map(items.cob if items is not None and "cob" in items else pd.Series(dtype=float)).fillna(0)
    it["sobre"] = np.where(it.directo, (it.contratado - it.proyectado).clip(lower=0), 0)
    it["venta"] = it.index.map(items.venta if items is not None and "venta" in items else pd.Series(dtype=float)).fillna(0)
    # pagado sin cobrar solo en ítems que se venden a AIRPLAN (los transversales sin venta no se cobran)
    it["sin_cobrar"] = np.where(it.directo & (it.venta > 0), (it.pagado - it.contratado * it.pc).clip(lower=0), 0)
    quien = C.groupby("cod").ctta.agg(lambda s: ", ".join(sorted(set(x[:28] for x in s))))
    it["contratistas"] = it.index.map(quien).fillna("")
    d = it[it.directo]
    out = dict(items=it)
    out["kpi"] = dict(contratado=d.contratado.sum(), pagado=d.pagado.sum(), pendiente=(d.contratado - d.pagado).clip(lower=0).sum(),
                      sobre=d.sobre.sum(), n_sobre=int((d.sobre > 0).sum()), sin_cobrar=d.sin_cobrar.sum(),
                      proy_contratado=d.loc[d.contratado > 0, "proyectado"].sum(),
                      ind_contratado=it[~it.directo].contratado.sum(), ind_pagado=it[~it.directo].pagado.sum(),
                      ult_acta=A.fecha.max(), n_contratos=C.contrato.nunique())
    out["frentes"] = d.groupby("fr").agg(contratado=("contratado", "sum"), pagado=("pagado", "sum"), proyectado=("proyectado", "sum"),
                                         sobre=("sobre", "sum"), sin_cobrar=("sin_cobrar", "sum"), cobrado=("cobrado", "sum"))
    # por contratista
    C2 = C.assign(directo=C.cod.isin(m.index))
    A2 = A.merge(C.drop_duplicates("contrato")[["contrato", "ctta"]], on="contrato", how="left")
    cc = C2.groupby("ctta").agg(contratos=("contrato", "nunique"), contratado=("base", "sum"))
    cc["pagado"] = A2.groupby("ctta").base.sum()
    cc["ult_acta"] = A2.groupby("ctta").fecha.max()
    ci = C2[C2.directo].groupby(["ctta", "cod"]).base.sum().rename("c").reset_index()
    ci["peso"] = ci.c / ci.cod.map(ci.groupby("cod").c.sum())
    ci["sobre"] = ci.peso * ci.cod.map(it.sobre)
    ci["sin_cobrar"] = ci.peso * ci.cod.map(it.sin_cobrar)
    cc = cc.join(ci.groupby("ctta")[["sobre", "sin_cobrar"]].sum()).fillna({"pagado": 0, "sobre": 0, "sin_cobrar": 0})
    cc["pct_pagado"] = np.where(cc.contratado > 0, cc.pagado / cc.contratado, np.nan)
    out["contratistas"] = cc.sort_values("contratado", ascending=False)
    # mes a mes: pagado contra lo contratado que ya se ejecutó (contratado del ítem × avance del frente en la curva S)
    Ad = A[A.cod.isin(m.index)].dropna(subset=["fecha"]).copy()
    Ad["mes"] = Ad.fecha.dt.to_period("M")
    pm = Ad.groupby("mes").base.sum()
    Pf = P[(P.fr != "TOTAL") & P.fr.notna()]
    ev = Pf.pivot_table(index="semana", columns="fr", values="EV", aggfunc="sum")
    cs = Pf.groupby("fr").costo.first()
    ok = [s for s in ev.index if s <= hoy and ev.loc[s].sum() > 0]
    av = (ev.loc[ok] / cs).fillna(0) if ok else pd.DataFrame()
    cfr = d.groupby("fr").contratado.sum()
    meses = pd.period_range(pm.index.min() if len(pm) else hoy.to_period("M"), hoy.to_period("M"), freq="M")
    filas = []
    for mm in meses:
        fin = min(mm.end_time.normalize(), hoy)
        sem = [s for s in ok if s <= fin]
        ej = (av.loc[sem[-1]].reindex(cfr.index).fillna(0) * cfr).sum() if sem else 0.0
        filas.append((mm, pm.get(mm, 0.0), ej))
    t = pd.DataFrame(filas, columns=["mes", "pagado", "ejecutado_acum"]).set_index("mes")
    t["pagado_acum"] = t.pagado.cumsum()
    t["pendiente_estimado"] = (t.ejecutado_acum - t.pagado_acum).clip(lower=0)
    out["mensual"] = t
    return out
